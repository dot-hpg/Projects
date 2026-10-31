"""Build the submission document (report/Target_SQL_Case_Study_Guntuku_Hariprasad.docx)
and convert it to PDF with LibreOffice. Run after run_queries, make_charts, render_tables."""
from datetime import date
from pathlib import Path
import os
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(Path(__file__).parent))
from query_catalog import Q  # noqa: E402
from insights import (ANALYSIS, ASSUMPTIONS, AUTHOR, EXEC_SUMMARY, HEADLINE_KPIS,  # noqa: E402
                      RECOMMENDATIONS, TITLE)

ROOT = Path(__file__).resolve().parents[1]
SHOTS, CHARTS, OUT = ROOT / "outputs" / "screenshots", ROOT / "charts", ROOT / "report"
STEM = "Target_SQL_Case_Study_Guntuku_Hariprasad"
BLUE = RGBColor(0x1C, 0x5C, 0xAB)
INK2 = RGBColor(0x52, 0x51, 0x4E)
REPO_URL = os.environ.get("REPO_URL", "https://github.com/dot-hpg/Projects/tree/main/target-brazil-sql-case-study")
TEXT_W = 7.0  # usable width in inches


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def cell_borders(cell, color="D9DDE3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def para(doc, text="", size=10, bold=False, italic=False, color=None, space_after=4, align=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
        if color is not None:
            r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def label(doc, text):
    p = para(doc, text, size=9.5, bold=True, color=BLUE, space_after=2)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.keep_with_next = True
    return p


def bullets(doc, items, size=9.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(1.5)


def code_block(doc, sql):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit"); tr_pr.append(cant)  # keep each query on one page
    c = t.rows[0].cells[0]
    shade(c, "F5F7FA"); cell_borders(c)
    c.width = Inches(TEXT_W)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lines = sql.rstrip().split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(7.5)
        if line.strip().startswith("--"):
            r.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
        if i < len(lines) - 1:
            r.add_break()


def picture(doc, path, max_w=TEXT_W, dpi=200, scale=0.82):
    w_px, _ = Image.open(path).size
    w_in = min(max_w, w_px / dpi * scale)
    doc.add_picture(str(path), width=Inches(w_in))
    doc.paragraphs[-1].paragraph_format.space_after = Pt(4)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.63))
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Title", 26)):
        s = doc.styles[name]
        s.font.name = "Calibri"; s.font.size = Pt(size); s.font.color.rgb = BLUE
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(6)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)
    # footer with page numbers
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run(f"{TITLE}  ·  {AUTHOR}  ·  page ")
    r.font.size = Pt(8); r.font.color.rgb = INK2
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = fp.add_run(); run.font.size = Pt(8); run.font.color.rgb = INK2
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)


def cover(doc):
    para(doc, "", space_after=60)
    para(doc, "BUSINESS CASE STUDY  ·  SQL", size=11, bold=True, color=INK2, space_after=6)
    p = para(doc, "Target Brazil", size=34, bold=True, color=BLUE, space_after=0)
    para(doc, "E-commerce operations analysis, 2016-2018", size=18, color=INK2, space_after=24)
    para(doc, "SQL analysis of 99,441 orders across customers, sellers, products, payments, freight, delivery "
              "and reviews, with insights and actionable recommendations.", size=11.5, space_after=30)
    t = doc.add_table(rows=1, cols=len(HEADLINE_KPIS))
    for i, (v, l) in enumerate(HEADLINE_KPIS):
        c = t.rows[0].cells[i]
        shade(c, "EEF3FB"); cell_borders(c, "FFFFFF")
        c.paragraphs[0].add_run(v).bold = True
        c.paragraphs[0].runs[0].font.size = Pt(15); c.paragraphs[0].runs[0].font.color.rgb = BLUE
        r = c.add_paragraph().add_run(l); r.font.size = Pt(8.5); r.font.color.rgb = INK2
    para(doc, "", space_after=40)
    for k, v in (("Author", AUTHOR), ("Context", "Scaler business case study"),
                 ("Tools", "Google BigQuery (Standard SQL) · DuckDB · Python (pandas, matplotlib)"),
                 ("Date", date.today().strftime("%d %B %Y")), ("Repository", REPO_URL)):
        p = para(doc, space_after=2)
        r = p.add_run(f"{k}:  "); r.bold = True; r.font.size = Pt(10.5)
        r = p.add_run(v); r.font.size = Pt(10.5)
    doc.add_page_break()


def front_matter(doc):
    doc.add_heading("Executive summary", level=1)
    bullets(doc, EXEC_SUMMARY, size=10)
    doc.add_heading("Top recommendations", level=2)
    for i, (h, body) in enumerate(RECOMMENDATIONS[:4], 1):
        p = para(doc, space_after=3)
        r = p.add_run(f"{i}. {h}. "); r.bold = True; r.font.size = Pt(10)
        r = p.add_run(body); r.font.size = Pt(10)
    para(doc, "All eight recommendations, with the evidence behind each, are in section 8.", size=9.5,
         italic=True, color=INK2)
    doc.add_heading("Dataset, approach & assumptions", level=2)
    para(doc, "Eight tables (orders, customers, order_items, payments, order_reviews, products, sellers, "
              "geolocation) were loaded into BigQuery as dataset target and mirrored locally in DuckDB. Every "
              "query below is BigQuery Standard SQL; the identical DuckDB version and the full results are in the "
              "repository. Result screenshots show the first 10 rows, as required.", size=9.5)
    bullets(doc, ASSUMPTIONS, size=9)
    picture(doc, ROOT / "data" / "schema.png", max_w=4.6, dpi=100, scale=1)
    para(doc, "Dataset schema: orders is the hub; items link to products and sellers; customers and sellers "
              "link to geolocation by zip prefix.", size=8.5, italic=True, color=INK2,
         align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_page_break()


def body(doc):
    current = None
    for item in Q:
        if item["section"] != current:
            if current is not None:
                doc.add_page_break()
            current = item["section"]
            doc.add_heading(current, level=1)
        h = doc.add_heading(f"Q{item['id']}  {item['title']}", level=2)
        h.paragraph_format.keep_with_next = True
        para(doc, item["question"], size=9.5, italic=True, color=INK2, space_after=3)
        label(doc, "Query")
        code_block(doc, item["bigquery"])
        label(doc, "Result (first 10 rows)")
        stem = f"{item['id'].replace('.', '_')}_{item['slug']}"
        picture(doc, SHOTS / f"{stem}.png")
        a = ANALYSIS[item["id"]]
        if a.get("chart"):
            picture(doc, CHARTS / a["chart"], max_w=6.2, dpi=170, scale=1)
        label(doc, "Insights")
        bullets(doc, a["insights"])
        label(doc, "Recommendations")
        bullets(doc, a["recs"])


def closing(doc):
    doc.add_page_break()
    doc.add_heading("8. Actionable insights & recommendations", level=1)
    para(doc, "Ranked by expected impact on customer satisfaction and revenue. Each one is backed by the queries "
              "referenced above.", size=10, color=INK2)
    for i, (h, b) in enumerate(RECOMMENDATIONS, 1):
        p = para(doc, space_after=1)
        r = p.add_run(f"{i}. {h}"); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
        para(doc, b, size=10, space_after=6)
    doc.add_heading("Suggested KPIs to track", level=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(("KPI", "Current value", "Target direction")):
        t.rows[0].cells[i].text = h
    for row in (("On-time delivery rate", "93.2% (6.8% late)", "≥ 96%, and ≤ 8% in every month"),
                ("Average delivery days", "12.1 (SP 8.3, RR 29.0)", "≤ 10 nationally; halve North gap"),
                ("Estimate padding (est. - actual)", "11.0 days early on average", "3-5 days, per state"),
                ("Same-state fulfilment share", "36.2% of items", "≥ 45%"),
                ("Freight as % of price (North/NE)", "20-29%", "≤ 18%"),
                ("Repeat-customer rate", "3.1%", "≥ 6%"),
                ("Average review score", "4.09", "≥ 4.3")):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    para(doc, "", space_after=8)
    para(doc, f"Code, data, notebook and full results: {REPO_URL}", size=9.5, italic=True, color=INK2)


def main():
    OUT.mkdir(exist_ok=True)
    doc = Document()
    setup(doc)
    cover(doc)
    front_matter(doc)
    body(doc)
    closing(doc)
    docx_path = OUT / f"{STEM}.docx"
    doc.save(docx_path)
    print("wrote", docx_path)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(docx_path)],
                       check=True, capture_output=True)
        print("wrote", OUT / f"{STEM}.pdf")


if __name__ == "__main__":
    main()
